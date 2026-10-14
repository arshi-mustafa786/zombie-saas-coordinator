import pptx

prs = pptx.Presentation(r"C:\Users\arshimustafa\Downloads\Submission Template _ AI Builder Cup.pptx")
print(f"Total slides: {len(prs.slides)}")
print(f"Slide width: {prs.slide_width / 914400} inches, height: {prs.slide_height / 914400} inches")

for idx, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {idx + 1} ====================")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            print(f"  Shape {s_idx} ({shape.name}): left={shape.left/914400:.2f}in, top={shape.top/914400:.2f}in, width={shape.width/914400:.2f}in, height={shape.height/914400:.2f}in")
            for p_idx, p in enumerate(shape.text_frame.paragraphs):
                text = p.text.strip()
                if text:
                    font_info = ""
                    if p.runs:
                        r0 = p.runs[0]
                        font_info = f"[font={r0.font.name}, size={r0.font.size.pt if r0.font.size else 'None'}pt, bold={r0.font.bold}]"
                    print(f"    P{p_idx}: {text} {font_info}")
