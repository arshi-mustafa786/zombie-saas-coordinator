import sys
import pptx

prs = pptx.Presentation(r"C:\Users\arshimustafa\Downloads\Submission Template _ AI Builder Cup.pptx")
print(f"Total slides: {len(prs.slides)}")

for idx, slide in enumerate(prs.slides):
    print(f"\n--- Slide {idx + 1} ---")
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                text = p.text.strip()
                if text:
                    print(f"  [Text]: {text}")
        elif shape.has_table:
            print("  [Table]:")
            for row in shape.table.rows:
                row_txt = [c.text.strip().replace('\n', ' ') for c in row.cells]
                print(f"    {' | '.join(row_txt)}")
