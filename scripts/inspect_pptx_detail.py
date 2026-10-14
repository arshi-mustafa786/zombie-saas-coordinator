import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE

prs = pptx.Presentation(r"C:\Users\arshimustafa\Downloads\Submission Template _ AI Builder Cup.pptx")
print(f"Slide width: {prs.slide_width}, height: {prs.slide_height}")

for idx, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {idx + 1} ====================")
    print(f"Layout name: {slide.slide_layout.name}")
    for s_idx, shape in enumerate(slide.shapes):
        shape_type = str(shape.shape_type)
        is_ph = shape.is_placeholder
        ph_fmt = f"(Placeholder: {shape.placeholder_format.type})" if is_ph else ""
        print(f"  Shape {s_idx}: name='{shape.name}', type={shape_type} {ph_fmt}, left={shape.left}, top={shape.top}, width={shape.width}, height={shape.height}")
        if shape.has_text_frame:
            for p_idx, p in enumerate(shape.text_frame.paragraphs):
                runs_text = "".join(r.text for r in p.runs)
                print(f"    P{p_idx}: '{p.text}' (runs: '{runs_text}')")
        if shape.has_table:
            print("    [Table]:")
            for r in shape.table.rows:
                print("      " + " | ".join(c.text.strip().replace('\n', ' ') for c in r.cells))
