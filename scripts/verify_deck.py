import pptx

prs = pptx.Presentation(r"c:\Users\arshimustafa\source\personal\Google AI Builder 2026\Zombie SAAS Coordinator\docs\Submission_AI_Builder_Cup_Zombie_SaaS_Consolidator.pptx")
print(f"Total slides: {len(prs.slides)}")

for idx, slide in enumerate(prs.slides):
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                txt = p.text.strip()
                if txt:
                    texts.append(txt)
    print(f"\n--- Slide {idx+1} ({len(slide.shapes)} shapes, {len(texts)} text paragraphs) ---")
    if texts:
        print(f"  Header: {texts[0]}")
        if len(texts) > 1:
            print(f"  Sample: {texts[1][:80]}...")
            print(f"  Total items: {len(texts)}")
